"""Focused tests for the host table reader's numeric fidelity."""
from __future__ import annotations

import math
from pathlib import Path
import struct
import tempfile
import unittest

from openpyxl import Workbook

from pyml_workbench.data import load_dataset, select_features_target


def _bits(value: float) -> int:
    return struct.unpack(">Q", struct.pack(">d", float(value)))[0]


class XlsxNumericFidelityTest(unittest.TestCase):
    def test_loader_preserves_binary64_sign_and_explicit_feature_selection(self) -> None:
        feature_ids = (
            "F-precision",
            "F-negative-zero",
            "F-nextafter",
            "F-subnormal",
            "F-max",
            "F-negative-max",
            "F-missing",
        )
        source_values = (
            1.2345678901234567,
            -0.0,
            math.nextafter(1.0, 2.0),
            float.fromhex("0x0.0000000000001p-1022"),
            float.fromhex("0x1.fffffffffffffp+1023"),
            -float.fromhex("0x1.fffffffffffffp+1023"),
            3.75,
        )

        with tempfile.TemporaryDirectory(prefix="pyml-xlsx-precision-") as temporary:
            workbook_path = Path(temporary) / "precision.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "FeatureMatrix"
            worksheet.append(["ObservationID", "Group", "AgeYears", *feature_ids])
            for observation_id, group, age_years, missing in (
                ("OBS-42", "control", 31, False),
                ("OBS-8", "case", 45, True),
            ):
                worksheet.append([observation_id, group, age_years, *([None] * len(feature_ids))])
                row_index = worksheet.max_row
                for offset, value in enumerate(source_values, start=4):
                    if missing and feature_ids[offset - 4] == "F-missing":
                        continue
                    cell = worksheet.cell(row=row_index, column=offset)
                    # Match the converter's numeric-cell encoding so -0.0 is
                    # represented in the XLSX XML as -0.0 rather than -0.
                    cell.value = repr(value)
                    cell.data_type = "n"
            workbook.save(workbook_path)

            dataset = load_dataset(workbook_path, sheet_name="FeatureMatrix")
            features, target = select_features_target(
                dataset,
                task="classification",
                target_column="Group",
                feature_columns=feature_ids,
            )

        self.assertEqual(list(features.columns), list(feature_ids))
        self.assertEqual(list(target), ["control", "case"])
        self.assertEqual(list(dataset.frame["ObservationID"]), ["OBS-42", "OBS-8"])
        self.assertTrue({"ObservationID", "Group", "AgeYears"}.isdisjoint(features.columns))

        actual = features.to_numpy(dtype="float64")
        self.assertFalse(math.isnan(float(actual[0, -1])))
        self.assertTrue(math.isnan(float(actual[1, -1])))
        for row_index in range(2):
            for column_index, expected in enumerate(source_values[:-1]):
                self.assertEqual(_bits(actual[row_index, column_index]), _bits(expected))


class CsvIdentityFidelityTest(unittest.TestCase):
    def test_identity_tokens_stay_text_without_changing_numeric_round_trip(self) -> None:
        with tempfile.TemporaryDirectory(prefix="pyml-csv-identity-") as temporary:
            csv_path = Path(temporary) / "identity.csv"
            csv_path.write_text(
                "sample_id,meta__group,meta__subject_id,feature__precise,"
                "feature__negative_zero,feature__integer,feature__missing\n"
                "001,1,NA,1.2345678901234567,-0.0,7,10\n"
                "NA,2,,0.10000000000000001,0.0,8,\n",
                encoding="utf-8",
            )

            frame = load_dataset(csv_path).frame

        self.assertEqual(frame["sample_id"].tolist(), ["001", "NA"])
        self.assertEqual(frame["meta__group"].tolist(), ["1", "2"])
        self.assertEqual(frame.loc[0, "meta__subject_id"], "NA")
        self.assertTrue(bool(frame["meta__subject_id"].isna().iloc[1]))

        self.assertEqual(frame["feature__precise"].dtype.kind, "f")
        self.assertIn(frame["feature__integer"].dtype.kind, ("i", "u"))
        self.assertTrue(bool(frame["feature__missing"].isna().iloc[1]))
        self.assertEqual(
            _bits(float(frame.loc[0, "feature__precise"])),
            _bits(1.2345678901234567),
        )
        self.assertEqual(
            _bits(float(frame.loc[1, "feature__precise"])),
            _bits(0.1),
        )
        self.assertEqual(
            _bits(float(frame.loc[0, "feature__negative_zero"])),
            _bits(-0.0),
        )


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
